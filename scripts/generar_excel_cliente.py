# generar_excel_cliente.py
# Simula un Excel "de verdad" hechos a mano en Excel: título combinado,
# nota antes de los encabezados y fechas con formato propio.
from pathlib import Path
from datetime import datetime

import openpyxl

carpeta = Path("datos_prueba/cliente")
carpeta.mkdir(parents=True, exist_ok=True)

wb = openpyxl.Workbook()
hoja = wb.active
hoja.title = "Ventas"

# Fila 1: título combinado (los clientes lo hacen SIEMPRE)
hoja.merge_cells("A1:C1")
hoja["A1"] = "Reporte de ventas - Cliente S.A."

# Fila 2: nota del área comercial
hoja["A2"] = "Generado el 24/10/2026 - Válido como respaldo fiscal"

# Fila 3: los encabezados REALES (para consolidar.py: --fila-encabezado 3)
encabezados = ["fecha", "producto", "monto"]
for col, valor in enumerate(encabezados, start=1):
    hoja.cell(row=3, column=col, value=valor)

# Filas 4 a 6: datos con fechas REALES de Excel (no texto) y formato propio
datos = [
    (datetime(2026, 1, 15), "Camisa", 15000),
    (datetime(2026, 2, 3), "Pantalón", 22000),
    (datetime(2026, 2, 20), "Zapatos", 48000),
]
for fila, (fecha, producto, monto) in enumerate(datos, start=4):
    hoja.cell(row=fila, column=1, value=fecha).number_format = "DD/MM/YYYY"
    hoja.cell(row=fila, column=2, value=producto)
    hoja.cell(row=fila, column=3, value=monto)

wb.save(carpeta / "ventas_cliente.xlsx")
print(f"✅ Excel 'de cliente' creado en {carpeta / 'ventas_cliente.xlsx'}")
